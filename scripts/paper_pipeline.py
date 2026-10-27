#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""paper-docx-optimizer 主流程脚本。

四个子命令串联起「解析 → 优化 → 定稿 → 交付」全链路：

    parse     结构解析：定位主论文/检测报告，抽取章节标题，建立工作目录与 manifest.json
    optimize  生成总揽与章节任务文档，执行第一轮规则化优化（可 --dry-run 预览）
    finalize  把优化稿另存为「最终版」，不覆盖原始 DOCX
    report    汇总解析、扫描、格式体检结果，输出交付报告

设计要点：
- 原始 DOCX 只读，所有产物写入 output_dir，绝不覆盖输入文件
- 文本回填走 run 级替换，保留字体、字号、加粗、上标等原有格式
- 术语 / 编号 / 数值守卫：改写若丢失 Latin 术语、图表公式编号或数字，自动回退该段
- docx2tex 为可选增强，环境缺失时自动降级为 python-docx 结构解析

依赖：python-docx（必需）、PyYAML（仅 ER 图脚本需要）
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx_ops import (  # noqa: E402
    build_spans,
    docx_integrity,
    dump_json,
    extract_headings,
    is_caption_style,
    is_skippable_style,
    load_config,
    mkdirs,
    now_iso,
    paragraph_plain_text,
    pick_docx,
    pick_report,
)

try:
    from docx import Document  # noqa: F401
except ImportError:  # pragma: no cover
    Document = None

SKILL_DIR = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = SKILL_DIR / "templates"

# ---------------------------------------------------------------- 章节优先级

# 关键词 → 优先级；命中即生效，未命中归入 P2
CHAPTER_PRIORITY_RULES: list[tuple[tuple[str, ...], str]] = [
    (("摘要", "abstract"), "P0"),
    (("总结与展望", "结论", "结语", "总结"), "P0"),
    (("相关技术", "相关工作", "文献综述", "研究现状", "理论基础"), "P1"),
    (("系统设计", "总体设计", "架构设计", "数据库设计"), "P1"),
    (("系统实现", "功能实现", "功能展示", "详细设计"), "P1"),
    (("测试", "实验", "评估", "结果分析"), "P1"),
    (("绪论", "引言", "前言"), "P2"),
    (("需求分析",), "P2"),
]

PRIORITY_LABEL = {"P0": "最高", "P1": "高", "P2": "常规"}

# 非优化对象：结构性章节，只做格式检查，不改内容
NON_TARGET_TITLES = (
    "参考文献", "致谢", "后记", "附录", "目录", "封面", "封底",
    "独创性声明", "原创性声明", "授权声明", "符号说明", "缩略词",
)

# ---------------------------------------------------------------- 规则词表

# T1：强 AI 标记，出现即可替换；T2：弱标记，仅在段落已被判定为高风险时替换
TIER1_REPLACEMENTS: dict[str, str] = {
    "这套系统": "本系统",
    "这个系统": "本系统",
    "打包在一起": "整合到统一平台",
    "有点麻烦": "存在不便",
    "都在网上完成": "整合为线上流程",
    "帮了大忙": "提供了便利",
    "一目了然": "便于识别",
    "非常醒目": "较为突出",
    "值得一提的是": "需要说明的是",
    "可以看出": "可见",
    "这说明": "这表明",
    "更像一个": "更接近",
    "从最终结果看": "从原型结果看",
    "整体来看": "总体来看",
    "从整体上看": "总体来看",
    "极大地提高了": "提高了",
    "显著提升了": "提升了",
    "显著提升": "提升",
    "有效提升了": "提升了",
    "有效提升": "提升",
    "有效地解决了": "缓解了",
    "充分体现了": "体现了",
    "为用户提供了便利": "降低了操作成本",
    "为用户提供便利": "降低操作成本",
    "具有重要的现实意义": "具备一定的应用价值",
    "奠定了坚实基础": "提供了基础",
    "提供了有力支撑": "提供了参考",
    "注入新动能": "带来改进",
    "值得注意的是，": "",
    "需要指出的是，": "",
}

TIER2_REPLACEMENTS: dict[str, str] = {
    "值得注意的是": "需要说明的是",
    "此外": "另外",
    "从而": "进而",
    "进而": "并",
    "与此同时": "同时",
    "综上所述": "综上",
    "总而言之": "总体而言",
    "在一定程度上": "部分",
    "较为": "相对",
    "十分": "较为",
    "非常": "较为",
}

REGEX_RULES: list[tuple[str, str]] = [
    # 三段式序列：同一句内的「首先…；其次…；最后…」压成并列短句
    (
        r"首先[，,]\s*([^。；;]{2,80})[；;]\s*其次[，,]\s*([^。；;]{2,80})[；;]\s*最后[，,]\s*([^。；;]{2,80})。",
        r"\1，\2，\3。",
    ),
    (
        r"首先[，,]\s*([^。；;]{2,80})[；;]\s*(?:其次|然后)[，,]\s*([^。；;]{2,80})。",
        r"\1，\2。",
    ),
    # 模板化背景：先处理带「近年来」前缀的长句，避免出现「近年来，近年来，」
    (r"(?:近年来|近几年|最近几年)[，,]?\s*随着(.{2,24}?)的(?:快速|迅猛|不断|持续)发展", r"近年来，\1持续发展"),
    (r"随着(.{2,24}?)的(?:快速|迅猛)发展", r"\1持续发展"),
    (r"随着(.{2,24}?)的不断(?:发展|普及|深入)", r"\1不断发展"),
    (r"“([^”]{1,12})”\s*这类问题", r"“\1”等问题"),
    (r"可以从(.+?)方面来谈。", r"主要体现在\1。"),
    (r"这是一条从", "该流程对应"),
    (r"为了说明", "用于说明"),
    (r"较快地完成", "完成"),
    (r"有助于强化", "有助于形成"),
    (r"本节主要(?:介绍|阐述)", "本节说明"),
    (r"(\S)。\1。", r"\1。"),
    (r"\s+", " "),
]

# 触发「规则改写」的 AI 模式（取自 ai_scan 模式库，保证口径一致）
TRIGGER_PATTERN_IDS = {"P01", "P02", "P03", "P04", "P05", "P06", "P07", "P08", "P09"}

# 术语 / 编号 / 数值守卫
LATIN_TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z0-9][A-Za-z0-9\.\-_/]*")
NUMBER_RE = re.compile(r"\d+(?:\.\d+)?%?")
LABEL_RE = re.compile(r"(?:图|表|式|公式|文献|第)\s?[\(\[]?\s?\d+(?:[\.\-－]\d+)*\s?[\)\]]?")


@dataclass
class Snippet:
    seq: str
    text: str
    risk: str


# ---------------------------------------------------------------- 报告解析

def parse_report(report_path: Path | None) -> list[Snippet]:
    """解析 AIGC / 查重报告，兼容 HTML 表格、Markdown 表格、项目符号三种写法。"""
    if not report_path or not report_path.exists():
        return []
    text = report_path.read_text(encoding="utf-8", errors="ignore")
    snippets: list[Snippet] = []

    for seq, body, risk in re.findall(
        r"<tr>\s*<td>\s*(\d+)\s*</td>\s*<td>(.*?)</td>\s*<td>\s*(高|中|低)\s*</td>",
        text,
        flags=re.S,
    ):
        snippets.append(Snippet(seq, _clean_cell(body), risk))

    for seq, body, risk in re.findall(
        r"^\|\s*(\d+)\s*\|\s*(.+?)\s*\|\s*(高|中|低)\s*\|", text, flags=re.M
    ):
        snippets.append(Snippet(seq, _clean_cell(body), risk))

    for risk, body in re.findall(r"^\s*[-*]\s*[\[【](高|中|低)[\]】]\s*(.+)$", text, flags=re.M):
        snippets.append(Snippet("", _clean_cell(body), risk))

    # 去重并限制体积
    seen = set()
    unique = []
    for item in snippets:
        key = (item.text[:40], item.risk)
        if key in seen or len(item.text) < 8:
            continue
        seen.add(key)
        unique.append(item)
    return unique


def _clean_cell(raw: str) -> str:
    text = re.sub(r"<[^>]+>", "", raw)
    text = text.replace("|", " ").replace("&nbsp;", " ")
    return re.sub(r"\s+", " ", text).strip()


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", "", text.replace(",", "，").replace(":", "："))


def match_snippet(text: str, snippets: list[Snippet]) -> str | None:
    norm = normalize_text(text)
    for item in snippets:
        sample = normalize_text(item.text)
        if len(sample) < 16:
            continue
        if sample in norm or norm in sample or sample[:28] in norm:
            return item.risk
    return None


# ---------------------------------------------------------------- 章节规划

def chapter_priority(title: str) -> str:
    low = title.lower()
    for keywords, priority in CHAPTER_PRIORITY_RULES:
        if any(kw in low for kw in keywords):
            return priority
    return "P2"


def parse_filters(value) -> list[str]:
    if not value:
        return []
    if isinstance(value, list):
        return [str(v).strip() for v in value if str(v).strip()]
    return [item.strip() for item in str(value).split(",") if item.strip()]


def chapter_enabled(title: str, allow: list[str], block: list[str]) -> bool:
    if allow and not any(token in title for token in allow):
        return False
    if block and any(token in title for token in block):
        return False
    return True


def is_non_target(title: str) -> bool:
    stripped = title.strip().replace(" ", "")
    return any(stripped.startswith(item) for item in NON_TARGET_TITLES)


def build_chapter_plan(spans, allow, block) -> list[dict]:
    plan = []
    for span in spans:
        if span["level"] not in (0, 1):
            continue
        non_target = is_non_target(span["title"])
        plan.append(
            {
                **span,
                "priority": chapter_priority(span["title"]),
                "kind": "matter" if non_target else "chapter",
                "enabled": (not non_target) and chapter_enabled(span["title"], allow, block),
            }
        )
    return plan


# ---------------------------------------------------------------- 改写

def apply_phrase_rules(text: str, aggressive: bool) -> str:
    updated = text
    for old, new in TIER1_REPLACEMENTS.items():
        updated = updated.replace(old, new)
    if aggressive:
        for old, new in TIER2_REPLACEMENTS.items():
            updated = updated.replace(old, new)
    for pattern, repl in REGEX_RULES:
        updated = re.sub(pattern, repl, updated)
    return updated.strip()


def rewrite_paragraph_runs(paragraph, aggressive: bool) -> tuple[str, bool]:
    """逐 run 应用替换规则，保留段内加粗 / 上标 / 字体差异。

    只在单个 run 内部替换，不改动 run 边界，因此不会破坏局部格式；
    代价是跨 run 的连续文本（如前半句在 run A、后半句在 run B）不会被规则命中，
    这部分留给精修环节处理。
    """
    changed = False
    for run in paragraph.runs:
        text = run.text
        if not text.strip():
            continue
        core = text.strip()
        lead = text[: len(text) - len(text.lstrip())]
        trail = text[len(text.rstrip()):]
        updated = apply_phrase_rules(core, aggressive)
        if updated != core:
            run.text = f"{lead}{updated}{trail}"
            changed = True
    return paragraph.text, changed


def restore_paragraph_runs(paragraph, snapshot: list[str]) -> None:
    for run, text in zip(paragraph.runs, snapshot):
        run.text = text


def guard_rewrite(old: str, new: str) -> list[str]:
    """返回新文本中丢失的关键 token；非空表示该段改写不可信。"""
    lost = []
    for pattern, label in (
        (LATIN_TOKEN_RE, "术语"),
        (LABEL_RE, "编号"),
        (NUMBER_RE, "数值"),
    ):
        before = set(pattern.findall(old))
        after = set(pattern.findall(new))
        missing = {t for t in before - after if t.strip()}
        if missing:
            lost.append(f"{label}丢失: {', '.join(sorted(missing)[:5])}")
    return lost


def ai_like_patterns(text: str) -> list[str]:
    """返回段落命中的 AI 模式名称（与 ai_scan 共用同一套正则）。"""
    try:
        from ai_scan import PATTERNS
    except Exception:  # pragma: no cover - ai_scan 缺失时降级为纯关键词判断
        return []
    return [name for pid, name, regex, _weight, _hint in PATTERNS
            if pid in TRIGGER_PATTERN_IDS and re.search(regex, text)]


def should_rewrite(paragraph, text: str, priority: str, report_risk: str | None) -> bool:
    if paragraph.style and is_caption_style(paragraph.style.name):
        return False
    if paragraph.style and is_skippable_style(paragraph.style.name):
        return False
    if len(text) < 12:
        return False
    if report_risk in {"高", "中"}:
        return True
    if priority == "P0":
        return True
    if priority == "P1":
        return bool(ai_like_patterns(text)) or any(
            word in text for word in ("首先", "其次", "最后", "从而", "此外", "值得注意的是")
        )
    return bool(ai_like_patterns(text))


# ---------------------------------------------------------------- 模板渲染

def chapter_stats(document, plan) -> dict[str, dict]:
    """按章节汇总 AI 模式命中情况（复用 ai_scan 的判据，保证口径一致）。"""
    try:
        from ai_scan import analyze_paragraph, load_lexicon, DEFAULT_LEXICON
    except Exception:  # pragma: no cover
        return {}
    lexicon = load_lexicon(DEFAULT_LEXICON)
    stats: dict[str, dict] = {}
    for item in plan:
        if item["kind"] != "chapter":
            continue
        weighted = 0.0
        chars = 0
        patterns: dict[str, int] = {}
        tier1: set[str] = set()
        for idx in range(item["start"], item["end"]):
            if idx >= len(document.paragraphs):
                break
            text = document.paragraphs[idx].text.strip()
            if len(text) < 12:
                continue
            result = analyze_paragraph(text, lexicon)
            weighted += result["score"] * result["chars"]
            chars += result["chars"]
            for pattern in result["patterns"]:
                patterns[pattern["name"]] = patterns.get(pattern["name"], 0) + 1
            tier1.update(result["tier1"])
        if not chars:
            continue
        stats[item["title"]] = {
            "score": round(weighted / chars, 1),
            "patterns": patterns,
            "tier1": sorted(tier1),
        }
    return stats


def render_template(name: str, values: dict[str, str]) -> str:
    path = TEMPLATE_DIR / name
    text = path.read_text(encoding="utf-8")
    for key, value in values.items():
        text = text.replace("{{" + key + "}}", str(value))
    return text


def write_overview(analysis_dir: Path, docx_path: Path, report_path, plan, snippets, allow, block) -> Path:
    matrix = []
    for item in plan:
        if item["kind"] == "matter":
            status = "跳过（非优化对象）"
        elif item["enabled"]:
            status = "启用"
        else:
            status = "跳过（章节过滤）"
        matrix.append(
            f"| {item['title']} | {item['priority']} | {PRIORITY_LABEL[item['priority']]} "
            f"| P{item['start']}–P{item['end'] - 1} | {status} |"
        )
    risk_by_priority: dict[str, list[str]] = {"P0": [], "P1": [], "P2": []}
    for item in plan:
        if item["kind"] == "chapter":
            risk_by_priority[item["priority"]].append(item["title"])

    targets = [item for item in plan if item["kind"] == "chapter"]
    matter = [item for item in plan if item["kind"] == "matter"]
    values = {
        "TIMESTAMP": now_iso(),
        "TITLE": docx_path.stem,
        "DOCX_PATH": docx_path,
        "REPORT_PATH": report_path if report_path else "未提供（按默认风险规则扫描）",
        "CHAPTER_COUNT": f"{len(targets)}（另有 {len(matter)} 个结构性章节不参与内容优化）",
        "STATUS": "已建立规划，待逐章改写",
        "P0_ITEMS": "、".join(risk_by_priority["P0"]) or "无",
        "P1_ITEMS": "、".join(risk_by_priority["P1"]) or "无",
        "P2_ITEMS": "、".join(risk_by_priority["P2"]) or "无",
        "CURRENT_STEP": "按优先级顺序处理 P0 章节",
        "NEXT_STEP": "生成章节任务 → 逐段改写 → 回填 DOCX → 评估",
        "BLOCKERS": "无" if snippets else "未提供检测报告，风险判定依据 references/risk-rules.md 默认规则",
        "CHAPTER_MATRIX": "\n".join(
            ["| 章节 | 优先级 | 说明 | 段落区间 | 状态 |", "| --- | --- | --- | --- | --- |"] + matrix
        ),
    }
    path = analysis_dir / "thesis_master_overview.md"
    path.write_text(render_template("master-overview-template.md", values), encoding="utf-8")

    extra = ["", "## 危险片段清单", ""]
    if snippets:
        for item in snippets[:40]:
            extra.append(f"- [{item.risk}] {item.text[:120]}")
    else:
        extra.append("- 未提供显式风险片段，按 references/risk-rules.md 的默认规则处理。")
    extra.append("")
    extra.append(f"- 白名单: {', '.join(allow) if allow else '未设置'}")
    extra.append(f"- 黑名单: {', '.join(block) if block else '未设置'}")
    path.write_text(path.read_text(encoding="utf-8") + "\n".join(extra), encoding="utf-8")
    return path


def write_chapter_tasks(
    analysis_dir: Path, plan, snippets, stats: dict | None = None, document=None
) -> list[str]:
    stats = stats or {}
    written = []
    order = 0
    for item in plan:
        if item["kind"] != "chapter":
            continue
        order += 1
        # 只认本章区间内的片段；别章的片段不借调，避免张冠李戴
        chapter_text = ""
        if document is not None:
            chapter_text = "".join(
                document.paragraphs[i].text
                for i in range(item["start"], min(item["end"], len(document.paragraphs)))
            )
        matched = [
            s
            for s in snippets
            if (chapter_text and s.text[:16] in chapter_text)
            or match_snippet(item["title"], [s])
        ]
        snippet_lines = "\n".join(f"- [{s.risk}] {s.text[:120]}" for s in matched[:6]) or "- 本章无报告片段"
        goals = {
            "P0": ("压缩「背景—方法—结果—意义」闭环", "拆掉并列式三点展望", "保留全部技术事实与数据"),
            "P1": ("清理教材式综述句", "打散高度对称的排比句", "术语、模块名、图表编号零改动"),
            "P2": ("剔除口语化表达", "收紧评价性措辞", "保持原文论证顺序"),
        }[item["priority"]]

        stat = stats.get(item["title"])
        if stat:
            pattern_text = "、".join(f"{name}×{count}" for name, count in
                                    sorted(stat["patterns"].items(), key=lambda kv: -kv[1])) or "未检出模式"
            ai_risks = f"章内风险分 {stat['score']}；命中 {pattern_text}"
            if stat["tier1"]:
                ai_risks += f"；强标记词：{'、'.join(stat['tier1'][:6])}"
        else:
            ai_risks = "按 references/ai-pattern-taxonomy.md 扫描"

        values = {
            "CHAPTER_ID": f"{order - 1:02d}",
            "CHAPTER_NAME": item["title"],
            "PARAGRAPH_RANGE": f"P{item['start']} – P{item['end'] - 1}",
            "PRIORITY": f"{item['priority']}（{PRIORITY_LABEL[item['priority']]}）",
            "STATUS": "待处理" if item["enabled"] else "跳过（被章节过滤排除）",
            "AI_RISKS": ai_risks,
            "PLAGIARISM_RISKS": "按 references/strategy-plagiarism.md 处理",
            "STYLE_RISKS": "按 references/strategy-polishing.md 处理",
            "GOAL_1": goals[0],
            "GOAL_2": goals[1],
            "GOAL_3": goals[2],
            "SOURCE_LOCATIONS": f"原文档第 {item['start']} 段起",
            "EDIT_NOTES": "待填写：逐段改写说明（句子级，禁止整段推翻）",
            "EVALUATION": "待填写：见 references/evaluation-criteria.md",
        }
        safe = re.sub(r"[^\w\u4e00-\u9fff]+", "_", item["title"]).strip("_") or "chapter"
        path = analysis_dir / f"chapter_{order - 1:02d}_{safe}.md"
        body = render_template("chapter-task-template.md", values)
        body += f"\n## 报告片段\n\n{snippet_lines}\n"
        path.write_text(body, encoding="utf-8")
        written.append(path.name)
    return written


# ---------------------------------------------------------------- 子命令

def ensure_manifest(input_dir: Path, output_dir: Path, args) -> dict:
    manifest_path = output_dir / "manifest.json"
    if manifest_path.exists() and not args.force:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if Path(manifest.get("docx_path", "")).exists():
            return manifest

    dirs = mkdirs(output_dir)
    docx_path = pick_docx(input_dir, args.docx_path).resolve()
    report_path = pick_report(input_dir, args.report_path)
    report_path = report_path.resolve() if report_path else None

    headings = extract_headings(docx_path)
    headings_path = dirs["analysis"] / "headings.json"
    dump_json(headings_path, headings)

    integrity = docx_integrity(docx_path)
    manifest = {
        "docx_path": str(docx_path),
        "report_path": str(report_path) if report_path else None,
        "input_dir": str(input_dir),
        "output_dir": str(output_dir),
        "headings_path": str(headings_path),
        "heading_count": len(headings),
        "docx_size_bytes": docx_path.stat().st_size if docx_path.exists() else 0,
        "docx_integrity_ok": integrity["ok"],
        "docx_integrity_errors": integrity["errors"],
        "created_at": now_iso(),
    }
    dump_json(manifest_path, manifest)

    structure = [
        "# DOCX 结构清单",
        "",
        f"- 主论文: `{docx_path}`",
        f"- 检测报告: `{report_path}`" if report_path else "- 检测报告: 未提供",
        f"- 标题条目数: {len(headings)}",
        f"- 包完整性: {'通过' if integrity['ok'] else '异常 → ' + '；'.join(integrity['errors'])}",
        "",
        "## 标题结构",
        "",
    ]
    if headings:
        for item in headings:
            indent = "  " * max(0, (item.get("level") or 1) - 1)
            structure.append(f"{indent}- P{item['index']} [{item['style']}] {item['text']}")
    else:
        structure.append("- 未识别到标题样式，请在 Word 中确认已套用「标题 1/2/3」或使用章节白名单显式指定")
    (dirs["analysis"] / "thesis_structure.md").write_text("\n".join(structure), encoding="utf-8")
    return manifest


def cmd_parse(args) -> int:
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    manifest = ensure_manifest(input_dir, output_dir, args)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


def cmd_optimize(args) -> int:
    if Document is None:
        print("缺少依赖 python-docx，请先执行: pip install python-docx", file=sys.stderr)
        return 2

    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    manifest = ensure_manifest(input_dir, output_dir, args)

    allow = parse_filters(args.chapter_allowlist)
    block = parse_filters(args.chapter_blocklist)
    docx_path = Path(manifest["docx_path"])
    report_path = Path(manifest["report_path"]) if manifest.get("report_path") else None

    headings = json.loads(Path(manifest["headings_path"]).read_text(encoding="utf-8"))
    snippets = parse_report(report_path)

    document = Document(str(docx_path))
    spans = build_spans(headings, len(document.paragraphs))
    plan = build_chapter_plan(spans, allow, block)

    analysis_dir = output_dir / "analysis"
    final_dir = output_dir / "final"
    analysis_dir.mkdir(parents=True, exist_ok=True)
    final_dir.mkdir(parents=True, exist_ok=True)

    overview_path = write_overview(analysis_dir, docx_path, report_path, plan, snippets, allow, block)
    chapter_files = write_chapter_tasks(
        analysis_dir, plan, snippets, chapter_stats(document, plan), document
    )

    priority_of: dict[int, str] = {}
    enabled_of: dict[int, bool] = {}
    for item in plan:
        for idx in range(item["start"], item["end"]):
            priority_of[idx] = item["priority"]
            enabled_of[idx] = item["enabled"]

    changes = []
    skipped = []
    for idx, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text or not enabled_of.get(idx, True):
            continue
        priority = priority_of.get(idx, "P2")
        report_risk = match_snippet(text, snippets)
        if not should_rewrite(paragraph, text, priority, report_risk):
            continue
        if args.dry_run:
            preview_text = apply_phrase_rules(text, aggressive=bool(report_risk in {"高", "中"} or priority == "P0"))
            if preview_text == text:
                continue
            lost = guard_rewrite(text, preview_text)
            if lost:
                skipped.append({"paragraph": idx, "reason": "；".join(lost), "text": text[:80]})
                continue
            changes.append(
                {
                    "paragraph": idx,
                    "priority": priority,
                    "mode": "dry-run",
                    "patterns": ai_like_patterns(text),
                    "before": text,
                    "after": preview_text,
                }
            )
            continue

        snapshot = [run.text for run in paragraph.runs]
        aggressive = bool(report_risk in {"高", "中"} or priority == "P0")
        new_text, changed = rewrite_paragraph_runs(paragraph, aggressive)
        if not changed or new_text.strip() == text:
            continue
        lost = guard_rewrite(text, new_text.strip())
        if lost:
            restore_paragraph_runs(paragraph, snapshot)
            skipped.append({"paragraph": idx, "reason": "；".join(lost), "text": text[:80]})
            continue
        changes.append(
            {
                "paragraph": idx,
                "priority": priority,
                "mode": "run-level",
                "patterns": ai_like_patterns(text),
                "before": text,
                "after": new_text.strip(),
            }
        )

    if args.dry_run:
        preview = ["# 改写预览（dry-run，未写入 DOCX）", ""]
        for item in changes:
            preview += [
                f"## P{item['paragraph']} · {item['priority']}",
                "",
                f"- 命中模式：{'、'.join(item['patterns']) or '—'}",
                f"- 原文：{item['before']}",
                f"- 改后：{item['after']}",
                "",
            ]
        if not changes:
            preview.append("- 本轮无可改写段落。")
        (analysis_dir / "rewrite_preview.md").write_text("\n".join(preview), encoding="utf-8")
        out_path = None
    else:
        out_path = final_dir / f"{docx_path.stem}_优化稿_v1.docx"
        document.save(str(out_path))

    result = {
        "docx_path": str(docx_path),
        "report_path": str(report_path) if report_path else None,
        "output_path": str(out_path) if out_path else None,
        "dry_run": bool(args.dry_run),
        "chapter_count": sum(1 for item in plan if item["kind"] == "chapter"),
        "matter_count": sum(1 for item in plan if item["kind"] == "matter"),
        "enabled_chapters": sum(1 for item in plan if item["enabled"]),
        "changed_paragraphs": len(changes),
        "rewrite_mode": "run-level（保留段内格式）",
        "skipped_paragraphs": len(skipped),
        "report_snippets": len(snippets),
        "chapter_allowlist": allow,
        "chapter_blocklist": block,
        "overview_path": str(overview_path),
        "chapter_files": chapter_files,
        "changes": changes,
        "skipped": skipped,
        "created_at": now_iso(),
    }
    dump_json(output_dir / "pipeline_result.json", result)

    summary = {k: v for k, v in result.items() if k not in {"changes", "skipped", "chapter_files"}}
    summary["chapter_files_count"] = len(chapter_files)
    print(json.dumps(summary, ensure_ascii=False, indent=2))

    if summary["chapter_count"] == 0 and not allow:
        print(
            "警告：未从标题样式中识别出任何章节，已退化为全文同优先级处理。\n"
            "      建议先跑 scripts/format_check.py --fix 规整标题样式，"
            "或用 --chapter-allowlist 显式圈定章节。",
            file=sys.stderr,
        )
    if not changes and not skipped:
        hint = "剩余内容请按 analysis/chapter_*.md 逐章精修"
        if not chapter_files:
            hint = "该文档缺少标题层级，只能做全文规则轮，后续精修需人工按段落处理"
        print(
            f"提示：本轮没有改写任何段落。规则轮只处理强标记词与模板句，{hint}。",
            file=sys.stderr,
        )
    return 0


def cmd_finalize(args) -> int:
    output_dir = Path(args.output_dir).resolve()
    result_path = output_dir / "pipeline_result.json"
    if not result_path.exists():
        print("未找到 pipeline_result.json，请先执行 optimize", file=sys.stderr)
        return 2
    result = json.loads(result_path.read_text(encoding="utf-8"))
    if not result.get("output_path"):
        print("上一次为 dry-run，未生成 DOCX；请先正式执行 optimize", file=sys.stderr)
        return 2

    source = Path(result["output_path"])
    if not source.exists():
        print(f"优化稿不存在: {source}", file=sys.stderr)
        return 2
    target = source.parent / source.name.replace("_优化稿_v1.docx", "_最终版.docx")
    shutil.copy2(source, target)

    integrity = docx_integrity(target)
    final = {**result, "final_output_path": str(target), "final_integrity_ok": integrity["ok"], "finalized_at": now_iso()}
    dump_json(output_dir / "finalize_result.json", final)
    print(json.dumps({k: final[k] for k in ("final_output_path", "final_integrity_ok")}, ensure_ascii=False, indent=2))
    return 0


def cmd_report(args) -> int:
    output_dir = Path(args.output_dir).resolve()
    analysis_dir = output_dir / "analysis"
    analysis_dir.mkdir(parents=True, exist_ok=True)

    def maybe(name: str):
        path = output_dir / name
        if not path.exists():
            path = analysis_dir / name
        if not path.exists():
            return None
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return None

    manifest = maybe("manifest.json") or {}
    pipeline = maybe("pipeline_result.json") or {}
    scan = maybe("ai_scan.json") or {}
    fmt = maybe("format_check.json") or {}

    lines = ["# 论文优化交付报告", "", f"- 生成时间: {now_iso()}", ""]
    lines += ["## 一、输入与结构", ""]
    lines += [
        f"- 原始 DOCX: `{manifest.get('docx_path', '未知')}`",
        f"- 检测报告: `{manifest.get('report_path')}`" if manifest.get("report_path") else "- 检测报告: 未提供",
        f"- 章节数: {manifest.get('heading_count', '未知')}",
        f"- 包完整性: {'通过' if manifest.get('docx_integrity_ok') else '未检测或异常'}",
        "",
    ]

    lines += ["## 二、内容改写", ""]
    if pipeline:
        lines += [
            f"- 处理模式: {'dry-run 预览' if pipeline.get('dry_run') else '正式写入'}",
            f"- 命中章节: {pipeline.get('enabled_chapters', 0)} / {pipeline.get('chapter_count', 0)}",
            f"- 改写段数: {pipeline.get('changed_paragraphs', 0)}（术语守卫回退 {pipeline.get('skipped_paragraphs', 0)} 段）",
            f"- 优化稿: `{pipeline.get('output_path')}`",
            "",
        ]
        if pipeline.get("changes"):
            lines += ["| 段落 | 优先级 | 改后片段 |", "| --- | --- | --- |"]
            for item in pipeline["changes"][:15]:
                lines.append(f"| P{item['paragraph']} | {item['priority']} | {item['after'][:60]} |")
            lines.append("")
    else:
        lines += ["- 尚未执行 optimize", ""]

    lines += ["## 三、AI 痕迹扫描", ""]
    if scan:
        sentence = scan.get("sentence") or {}
        cv = sentence.get("cv")
        verdict = "偏均匀（更像机器文本）" if isinstance(cv, (int, float)) and cv < 0.28 else "分布正常"
        lines += [
            f"- 全文风险分: {scan.get('overall_score', '未知')} / 100",
            f"- 句长标准差: {sentence.get('std', '未知')}（均值 {sentence.get('mean', '未知')}，变异系数 {cv if cv is not None else '未知'} → {verdict}）",
            f"- 高风险段落: {scan.get('high_risk_count', 0)}，中风险 {scan.get('medium_risk_count', 0)}",
            f"- 短句占比: {sentence.get('short_ratio', '未知')}；长句占比: {sentence.get('long_ratio', '未知')}",
            "",
        ]
    else:
        lines += ["- 尚未执行 ai_scan", ""]

    lines += ["## 四、格式体检", ""]
    if fmt:
        lines += [
            f"- 合规分: {fmt.get('score', '未知')} / 100",
            f"- 问题项: 严格 {fmt.get('errors', 0)} / 提醒 {fmt.get('warnings', 0)}",
            "",
        ]
        for item in (fmt.get("findings") or [])[:15]:
            lines.append(f"- [{item['level']}] {item['item']}：{item['detail']}")
        lines.append("")
    else:
        lines += ["- 尚未执行 format_check", ""]

    final_dir = output_dir / "final"
    final_docs = sorted(final_dir.glob("*.docx")) if final_dir.exists() else []
    lines += [
        "## 五、交付清单",
        "",
        f"- {pipeline.get('output_path') or '(待生成) 优化稿'}",
        "- analysis/thesis_master_overview.md、analysis/chapter_*.md",
        "- analysis/ai_scan.md、analysis/format_check.md（若已执行对应检查）",
    ]
    for doc in final_docs:
        lines.append(f"- {doc}（{doc.stat().st_size // 1024} KB）")
    lines += [
        "",
        f"> 共 {len(final_docs)} 份 docx 产物；请以「最终版」为交付对象，「优化稿_v1」保留用于对比回退。",
        "> 提醒：降 AI / 降重只做语言层调整，事实、数据、结论与引用关系必须由作者确认。",
        "",
    ]
    path = analysis_dir / "delivery_report.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"report": str(path)}, ensure_ascii=False, indent=2))
    return 0


# ---------------------------------------------------------------- 入口

def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="论文 DOCX 一体化优化流水线")
    sub = parser.add_subparsers(dest="command", required=True)

    def add_common(p):
        p.add_argument("--config", help="JSON 配置文件（与命令行参数同名键）")
        p.add_argument("--input-dir", help="输入目录（含论文 docx 与可选检测报告）")
        p.add_argument("--output-dir", help="输出目录")
        p.add_argument("--docx-path", help="显式指定主论文 docx")
        p.add_argument("--report-path", help="显式指定检测报告")
        p.add_argument("--force", action="store_true", help="忽略已有 manifest.json，重新解析")

    p_parse = sub.add_parser("parse", help="结构解析并建立工作目录")
    add_common(p_parse)
    p_parse.set_defaults(func=cmd_parse)

    p_opt = sub.add_parser("optimize", help="生成规划文档并执行第一轮规则化优化")
    add_common(p_opt)
    p_opt.add_argument("--chapter-allowlist", help="只处理命中的章节（逗号分隔，包含匹配）")
    p_opt.add_argument("--chapter-blocklist", help="跳过命中的章节（逗号分隔，包含匹配）")
    p_opt.add_argument("--dry-run", action="store_true", help="只输出改写预览，不改 DOCX")
    p_opt.set_defaults(func=cmd_optimize)

    p_fin = sub.add_parser("finalize", help="另存为最终版并校验完整性")
    p_fin.add_argument("--output-dir", required=True)
    p_fin.set_defaults(func=cmd_finalize)

    p_rep = sub.add_parser("report", help="汇总各环节结果，输出交付报告")
    p_rep.add_argument("--output-dir", required=True)
    p_rep.set_defaults(func=cmd_report)

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    if getattr(args, "config", None):
        config = load_config(args.config)
        for key in ("input_dir", "output_dir", "docx_path", "report_path",
                    "chapter_allowlist", "chapter_blocklist"):
            if getattr(args, key, None) in (None, False) and config.get(key):
                setattr(args, key, config[key])
        if config.get("dry_run") and hasattr(args, "dry_run"):
            args.dry_run = True
    if hasattr(args, "input_dir") and hasattr(args, "output_dir"):
        if not args.input_dir or not args.output_dir:
            parser.error("input-dir 与 output-dir 为必填（或通过 --config 提供）")
    try:
        return args.func(args)
    except FileNotFoundError as exc:
        print(f"找不到文件：{exc}", file=sys.stderr)
        print("请确认 --input-dir 指向包含论文 docx 的目录，必要时用 --docx-path 显式指定。", file=sys.stderr)
        return 2
    except ValueError as exc:
        print(f"参数或数据有误：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
